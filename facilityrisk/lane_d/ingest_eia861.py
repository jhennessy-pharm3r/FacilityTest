"""Stage real EIA-861 reliability, state totals and service-territory source rows.

Source identities are utility/report observations, NOT verified insured properties.
Original ZIP must be retained separately. Python stdlib + openpyxl for official XLSX.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import sqlite3
import urllib.request
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

SOURCE_PAGE = "https://www.eia.gov/electricity/data/eia861/"
SOURCE = SOURCE_PAGE + "zip/f861{year}.zip"
NULLS = {"", ".", "na", "n/a", "nan", "none", "null", "--"}
UTILITY_ID_HEADERS = {"utilitynumber", "utilityid", "utilitynumberid"}
MAX_MEMBER_BYTES = 40_000_000


def canon(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def cell(value):
    return "" if value is None else str(value).strip()


def sha256(blob):
    return hashlib.sha256(blob).hexdigest()


def row_digest(cells):
    return sha256(json.dumps(cells, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))


def forward_fill(cells, width):
    result, last = [], ""
    for i in range(width):
        value = cell(cells[i]) if i < len(cells) else ""
        if value and value != ".":
            last = value
        result.append(last)
    return result


def table_rows(member, data):
    """Keep every original worksheet distinct; do not concatenate workbook sheets."""
    if member.lower().endswith(".csv"):
        yield "CSV", list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    elif member.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for ws in wb:
                yield ws.title, [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        raise ValueError("Unsupported source: " + member)


def column_schema(rows, header_at):
    labels = [cell(v) for v in rows[header_at]]
    width = len(labels)
    groups = forward_fill(rows[0], width) if header_at >= 2 else [""] * width
    conditions = forward_fill(rows[header_at - 1], width) if header_at >= 2 else [""] * width
    return [
        {"column_1based": i + 1, "group": groups[i], "condition": conditions[i],
         "label": label or "_unnamed_" + str(i + 1)}
        for i, label in enumerate(labels)
    ]


def find_header(rows, kind):
    for i, row in enumerate(rows[:25]):
        cols = [canon(v) for v in row]
        has_id = bool(UTILITY_ID_HEADERS.intersection(cols))
        has_state = "state" in cols
        if kind == "state_total" and not has_id and has_state and "datayear" in cols:
            return i
        if kind in {"reliability", "service_territory"} and has_id and has_state:
            if kind == "reliability" and not any("saidi" in c or "saifi" in c for c in cols):
                continue
            return i
    return None


def process_sheet(member, sheet, rows, year, kind):
    header_at = find_header(rows, kind)
    stats = {
        "source_member": member, "source_sheet": sheet, "source_kind": kind,
        "physical_row_count": len(rows), "header_row_1based": None,
        "nonempty_data_rows": 0, "blank_data_rows": 0,
        "accepted": 0, "rejected": 0, "exact_duplicates": 0,
        "rows_without_reliability_metric": 0, "columns": [],
    }
    if header_at is None:
        stats["skipped_reason"] = "no_matching_header"
        return [], [], stats
    stats["header_row_1based"] = header_at + 1
    schema = column_schema(rows, header_at)
    stats["columns"] = schema
    headers = [canon(c["label"]) for c in schema]
    id_idx = next((i for i, key in enumerate(headers) if key in UTILITY_ID_HEADERS), None)
    state_idx = headers.index("state")
    year_idx = headers.index("datayear") if "datayear" in headers else None
    name_idx = next((i for i, name in enumerate(headers) if name in {"utilityname", "utilitycompany"}), None)
    county_idx = next((i for i, name in enumerate(headers) if "county" in name), None)
    seen, accepted, rejected = set(), [], []
    for line, source_row in enumerate(rows[header_at + 1:], start=header_at + 2):
        cells = [cell(v) for v in source_row]
        if not any(cells):
            stats["blank_data_rows"] += 1
            continue
        stats["nonempty_data_rows"] += 1
        def get(index):
            return cells[index] if index is not None and index < len(cells) else ""
        uid = get(id_idx)
        state = get(state_idx)
        data_year = get(year_idx)
        digest = row_digest(cells)
        reason = None
        if year_idx is not None and data_year != str(year):
            reason = "invalid_reporting_year"
        elif not re.fullmatch(r"[A-Z]{2}", state.upper()):
            reason = "invalid_state"
        elif kind != "state_total" and not re.fullmatch(r"[0-9]+", uid):
            reason = "invalid_utility_id"
        elif digest in seen:
            reason = "duplicate_exact_row"
            stats["exact_duplicates"] += 1
        if reason:
            rejected.append({"source_member": member, "source_sheet": sheet,
                             "source_sheet_row": line, "reason": reason,
                             "row_sha256": digest, "raw_cells": cells})
            continue
        seen.add(digest)
        measurements = [
            get(i) for i, c in enumerate(headers)
            if any(metric in c for metric in ("saidi", "saifi", "caidi"))
        ]
        no_metrics = kind == "reliability" and not any(v.lower() not in NULLS for v in measurements)
        stats["rows_without_reliability_metric"] += int(no_metrics)
        # This is an ingest-generated identifier, not a source-native facility/event ID.
        record_key = sha256(json.dumps([year, member, sheet, line, digest], separators=(",", ":")).encode())
        accepted.append({
            "source": "EIA861", "source_kind": kind,
            "source_year": year, "source_member": member, "source_sheet": sheet,
            "source_sheet_row": line, "source_utility_id": uid or None,
            "reported_data_year": data_year or None, "reported_state": state,
            "reported_utility_name": get(name_idx) or None,
            "reported_county": get(county_idx) or None,
            "row_sha256": digest, "ingestion_row_key": record_key,
            "missing_reliability_metrics": bool(no_metrics),
            "raw_cells": cells,
            "facility_match_status": "NOT_ASSESSED",
            "evidence_type": {
                "reliability": "UTILITY_AGGREGATE_CONTEXT",
                "state_total": "STATE_AGGREGATE_CONTEXT",
                "service_territory": "COUNTY_SERVICE_TERRITORY_CANDIDATE",
            }[kind],
        })
    stats["accepted"], stats["rejected"] = len(accepted), len(rejected)
    if len(accepted) + len(rejected) != stats["nonempty_data_rows"]:
        raise RuntimeError("Data-row reconciliation failed for " + member + "#" + sheet)
    return accepted, rejected, stats


def normalize(member, contents, year):
    """Backward compatible utility-only API used by offline fixtures."""
    accepted, rejected, duplicates, null_metrics = [], [], 0, 0
    for sheet, rows in table_rows(member, contents):
        kind = "state_total" if "total" in sheet.lower() else "reliability"
        found, bad, stats = process_sheet(member, sheet, rows, year, kind)
        accepted += [r for r in found if kind == "reliability"]
        rejected += bad
        duplicates += stats["exact_duplicates"]
        null_metrics += stats["rows_without_reliability_metric"]
    return accepted, rejected, duplicates, null_metrics


def write_jsonl(path, objects):
    with Path(path).open("w", encoding="utf-8") as f:
        for record in objects:
            f.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")


def build_local_index(output, datasets):
    """Review-only staging SQLite. Never imply durable production/deployment."""
    target = output / "lane_d_eia861_staging.sqlite"
    con = sqlite3.connect(target)
    try:
        con.execute("""CREATE TABLE IF NOT EXISTS source_observations (
            ingestion_row_key TEXT PRIMARY KEY,
            source_kind TEXT NOT NULL,
            reporting_year INTEGER NOT NULL,
            utility_id TEXT,
            state TEXT,
            county TEXT,
            source_member TEXT NOT NULL,
            source_sheet TEXT NOT NULL,
            source_row INTEGER NOT NULL,
            row_sha256 TEXT NOT NULL,
            original_row_json TEXT NOT NULL
        )""")
        con.execute("CREATE INDEX IF NOT EXISTS idx_lane_d_utility ON source_observations (reporting_year, utility_id, state)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_lane_d_geography ON source_observations (source_kind, state, county)")
        for rows in datasets.values():
            con.executemany(
                "INSERT INTO source_observations VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                [(r["ingestion_row_key"], r["source_kind"], r["source_year"], r["source_utility_id"],
                  r["reported_state"], r["reported_county"], r["source_member"], r["source_sheet"],
                  r["source_sheet_row"], r["row_sha256"], json.dumps(r["raw_cells"], ensure_ascii=False))
                 for r in rows],
            )
        con.commit()
        if con.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise RuntimeError("Staging SQLite integrity check failed")
        count = con.execute("SELECT COUNT(*) FROM source_observations").fetchone()[0]
        expected = sum(len(v) for v in datasets.values())
        if count != expected:
            raise RuntimeError(f"Staging index rows={count}, source rows={expected}")
        totals = dict(con.execute("SELECT source_kind, COUNT(*) FROM source_observations GROUP BY source_kind"))
        return {"stage": "INDEXED_LOCAL_STAGING", "records": count, "by_source_kind": totals,
                "database_path": str(target), "integrity_check": "ok", "durable_backend_records": 0}
    finally:
        con.close()


def run(year, output, archive=None):
    if not 2013 <= year <= datetime.now(timezone.utc).year:
        raise ValueError("Outside EIA-861 reliability collection period")
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    source_url = SOURCE.format(year=year)
    if archive:
        archive_path = Path(archive)
        if not archive_path.is_file():
            raise FileNotFoundError(archive)
    else:
        archive_path = output / ("f861%d.zip" % year)
        temp = archive_path.with_suffix(".zip.partial")
        try:
            with urllib.request.urlopen(
                urllib.request.Request(source_url, headers={"User-Agent": "FacilityRiskLaneD/1.0"}),
                timeout=90,
            ) as response, temp.open("wb") as f:
                shutil.copyfileobj(response, f)
            temp.replace(archive_path)
        finally:
            temp.unlink(missing_ok=True)
    if not zipfile.is_zipfile(archive_path):
        raise ValueError("EIA source archive is not ZIP")
    archive_hash = sha256(archive_path.read_bytes())
    datasets = {"reliability": [], "state_total": [], "service_territory": []}
    rejected, sheets, members = [], [], []
    with zipfile.ZipFile(archive_path) as z:
        for member in sorted(z.namelist()):
            filename = Path(member).name.lower()
            if "reliability" in filename:
                member_kind = "reliability"
            elif "service_territory" in filename or "service territory" in filename:
                member_kind = "service_territory"
            else:
                continue
            if not filename.endswith((".xlsx", ".xlsm", ".csv")):
                continue
            info = z.getinfo(member)
            if info.file_size > MAX_MEMBER_BYTES:
                raise ValueError("Oversized EIA source workbook: " + member)
            data = z.read(member)
            members.append({"path": member, "sha256": sha256(data), "bytes": len(data)})
            for sheet, original_rows in table_rows(member, data):
                kind = "state_total" if member_kind == "reliability" and "total" in sheet.lower() else member_kind
                rows, bad, stats = process_sheet(member, sheet, original_rows, year, kind)
                if kind == "reliability" and stats.get("skipped_reason"):
                    raise ValueError("Reliability sheet missing recognizable utility + metric header: " + sheet)
                datasets[kind].extend(rows)
                rejected.extend(bad)
                sheets.append(stats)
    if not any(s["source_kind"] == "reliability" for s in sheets):
        raise ValueError("No EIA-861 reliability workbook found; inspect ZIP layout")
    if not datasets["reliability"]:
        raise ValueError("No source-validated utility reliability rows; refusing empty import")
    for kind, records in datasets.items():
        write_jsonl(output / ("eia861_" + kind + ".jsonl"), records)
    write_jsonl(output / "eia861_rejected.jsonl", rejected)
    index = build_local_index(output, datasets)
    counts = {key: len(value) for key, value in datasets.items()}
    manifest = {
        "source": "EIA861", "publisher": "U.S. Energy Information Administration",
        "source_page": SOURCE_PAGE, "official_archive_url": source_url,
        "reporting_year": year, "run_utc": datetime.now(timezone.utc).isoformat(),
        "source_origin": "local_provided_archive" if archive else "official_live_download",
        "archive_sha256": archive_hash, "archive_bytes": archive_path.stat().st_size,
        "raw_archive_path": str(archive_path), "members": members, "sheets": sheets,
        "accepted_by_source_kind": counts, "accepted": sum(counts.values()),
        "rejected": len(rejected), "rejection_reasons": dict(Counter(r["reason"] for r in rejected)),
        "unique_reliability_utility_ids": len({r["source_utility_id"] for r in datasets["reliability"]}),
        "unique_territory_utility_ids": len({r["source_utility_id"] for r in datasets["service_territory"]}),
        "rows_without_reliability_metric": sum(r["missing_reliability_metrics"] for r in datasets["reliability"]),
        "local_index": index, "index_status": "LOCAL_STAGING_ONLY", "durable_backend_indexed": 0,
        "verified_facility_matches": 0, "deployed_rows": 0,
        "evidence_state": "SOURCE_CONTENT_CHECKED_NOT_SITE_VERIFIED",
        "limitations": [
            "Utility IDs identify survey respondents, not insured properties.",
            "Repeated SAIDI/SAIFI/CAIDI headers require column index, standard and MED condition.",
            "County territory means reported distribution presence, not exclusive service or a property connection.",
            "Aggregates may reflect different IEEE/other methods; no direct rankings without normalization.",
            "No derived facility outage, downtime, financial loss or insurance coverage claim.",
            "GitHub Actions artifact/local SQLite are ephemeral staging, not a durable production service.",
        ],
    }
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return manifest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--zip", dest="archive", type=Path)
    args = parser.parse_args()
    print(json.dumps(run(args.year, args.output, args.archive), indent=2))


if __name__ == "__main__":
    main()
