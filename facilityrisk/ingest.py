"""Ingest genuine curated evidence from the analyst's source workbook.

The underlying workbook must be supplied separately. The imported database is
ignored by Git, since the analyst's evaluation may be proprietary.
"""
from __future__ import annotations
import hashlib
import json
import re
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from .db import add_run, complete_run, initialize

SITE_SHEETS = [
 'Pfizer Rocky Mount', 'Pfizer Groton Historical', 'Baxter North Cove',
 'Amgen Juncos', 'Divis Chippada Unit II', 'Aurobindo Pydibhimavaram XI',
 'Dr Reddy Miryalaguda V', 'Boston Scientific Galway',
 'Abbott Clonmel Vascular', 'Baxter Vantive Mountain Home',
 'Boston Scientific Spencer', 'Pfizer Kalamazoo Portage'
]
SEC_SITE_LINKS = {
 'BAX': 'Baxter North Cove',
 'PFE': 'Pfizer Rocky Mount',
 'AMGN': 'Amgen Juncos',
 'RDY': 'Dr Reddy Miryalaguda V'
}
FIELD_PREFIXES = ['SITE', 'ADDRESS', 'IDENTIFIERS', 'CROSS-DATASET FINDING', 'COVERAGE', 'GAPS / CAUTIONS']
ORIGIN = 'Facility Risk Data Exploration.xlsx / analyst research, reviewed 2026-10-20'


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def stable_id(*parts) -> str:
    raw = json.dumps(parts, ensure_ascii=False, default=str, separators=(',', ':'))
    return hashlib.sha256(raw.encode('utf-8')).hexdigest()[:24]


def slug(s: str) -> str:
    return re.sub(r'-+', '-', re.sub(r'[^a-z0-9]+', '-', s.lower())).strip('-')


def site_id(tab: str) -> str:
    return 'site-' + slug(tab)


def cell_text(value) -> str:
    return str(value).strip() if value is not None else ''


def suffix(line: str) -> str:
    return line.split('|', 1)[1].strip() if '|' in line else line.strip()


def risk_domain(source: str, program: str = '') -> str:
    k = (source + ' ' + program).lower()
    if 'aqueduct' in k:
        return 'Water stress'
    if source.lower().startswith('sec'):
        return 'Financial / supply consequences'
    if source.lower() == 'fda':
        return 'Manufacturing regulation'
    if 'osha' in k:
        return 'Workplace safety'
    if any(x in k for x in ['fema','noaa','usgs']):
        return 'Natural hazards'
    if any(x in k for x in ['epa','environmental','tri','npdes','rcra']):
        return 'Environmental compliance'
    return 'Facility identity / operations'


def parse_site_metadata(sheet) -> dict:
    rows = [cell_text(sheet.cell(i,1).value) for i in range(1,9)]
    idline = rows[3]
    fei_m = re.search(r'\bFEI:\s*(\d+)', idline)
    cc_m = re.search(r'Country:\s*(.*?)\s*\|\s*Site point', idline)
    ll_m = re.search(r'\(lat/lon\):\s*(-?\d+(?:\.\d+)?)\s*/\s*(-?\d+(?:\.\d+)?)', idline)
    case_m = re.search(r'\|\s*Case:\s*(.*)$', idline)
    site = suffix(rows[1])
    parts = [p.strip() for p in site.split('  \u2022  ', 1)]
    coords = (float(ll_m.group(1)),float(ll_m.group(2))) if ll_m else (None,None)
    if coords[0] is not None and not(-90<=coords[0]<=90 and -180<=coords[1]<=180):
        coords = (None,None)
    return dict(id=site_id(sheet.title), name=rows[0].title(),
                company=parts[0] if parts else site, activity=parts[1] if len(parts)>1 else '',
                address=suffix(rows[2]),country=cc_m.group(1).strip() if cc_m else '',
                fei=fei_m.group(1) if fei_m else None,latitude=coords[0],longitude=coords[1],
                site_case=case_m.group(1).strip() if case_m else '',
                research_summary=suffix(rows[4]),source_coverage_note=suffix(rows[5]),
                data_gaps=suffix(rows[6]))


def ingest_workbook(conn, workbook_path: str | Path) -> dict:
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    w = load_workbook(path, read_only=True, data_only=True)
    missing = [s for s in [*SITE_SHEETS, 'SEC Findings', 'Public Evidence for SEC FInding', 'Files Summary', 'Field Information'] if s not in w.sheetnames]
    if missing:
        raise ValueError(f'Missing expected worksheet(s): {missing}')
    initialize(conn)
    run_id = uuid.uuid4().hex
    now = utcnow()
    add_run(conn,run_id,'analyst_workbook',now,sha)
    summary = Counter()
    try:
        # Replace workbook-owned rows atomically so changes/removals do not leave stale evidence.
        # Facility IDs remain stable; live-feed evidence and snapshots are preserved.
        for table in ('evidence', 'sec_findings', 'public_evidence', 'sources', 'field_catalog'):
            conn.execute(f'DELETE FROM {table} WHERE origin=?', (ORIGIN,))
        for tab in SITE_SHEETS:
            sh = w[tab]
            meta = parse_site_metadata(sh)
            conn.execute('''INSERT INTO facilities
            (id,name,company,activity,address,country,fei,latitude,longitude,site_case,
             research_summary,source_coverage_note,data_gaps,origin,last_imported)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(id) DO UPDATE SET
            name=excluded.name,company=excluded.company,activity=excluded.activity,
            address=excluded.address,country=excluded.country,fei=excluded.fei,
            latitude=excluded.latitude,longitude=excluded.longitude,
            site_case=excluded.site_case,research_summary=excluded.research_summary,
            source_coverage_note=excluded.source_coverage_note,data_gaps=excluded.data_gaps,
            origin=excluded.origin,last_imported=excluded.last_imported''',
            (*[meta[k] for k in ['id','name','company','activity','address','country','fei','latitude','longitude','site_case','research_summary','source_coverage_note','data_gaps']],ORIGIN,now))
            summary['facilities'] += 1
            found_header=False
            for row in sh.iter_rows(values_only=True):
                if not found_header:
                    if len(row)>=11 and cell_text(row[0])=='Source' and cell_text(row[1])=='Dataset / program':
                        found_header=True
                    continue
                vals = [cell_text(v) for v in (list(row[:11])+[None]*11)[:11]]
                if not vals[0]:
                    continue
                if vals[0] in {'Source','CURATED EVIDENCE','NOTES'}:
                    continue
                evidence_id = stable_id('workbook',tab, *vals)
                conn.execute('''INSERT INTO evidence
                (id,facility_id,source,program,date_raw,record_type,finding,value_text,
                 record_ref,evidence_scope,match_quality,source_url,limitations,risk_domain,
                 origin,availability_date,spatial_distance_km,retrieval_at,last_imported)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(id) DO UPDATE SET last_imported=excluded.last_imported''',
                (evidence_id,meta['id'], *vals[:9], vals[9], vals[10],
                 risk_domain(vals[0],vals[1]),ORIGIN,None,None,None,now))
                summary['evidence'] += 1
            if not found_header:
                raise ValueError(f'Evidence header missing in {tab}')
        findings = list(w['SEC Findings'].values)
        sec_header = next((i for i,row in enumerate(findings) if row[0]=='Finding ID'),None)
        if sec_header is None: raise ValueError('SEC findings header not found')
        for row in findings[sec_header+1:]:
            vals = [cell_text(x) for x in (list(row[:20])+[None]*20)[:20]]
            if not vals[0]: continue
            group = vals[0].split('-',1)[0]
            conn.execute('''INSERT INTO sec_findings
            (id,company,facility_label,finding_type,incident_window,finding,supply_impact,
             financial_impact,insurance_recovery,prior_public_hazard,attribution,country,
             address,filing_form,reporting_period,filing_date,section,filing_url,
             public_evidence_ids,caveats,facility_id,event_group,origin)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(id) DO UPDATE SET
            financial_impact=excluded.financial_impact,caveats=excluded.caveats,
            facility_id=excluded.facility_id,origin=excluded.origin''',
            (*vals,site_id(SEC_SITE_LINKS[group]) if group in SEC_SITE_LINKS else None,group,ORIGIN))
            summary['sec_findings'] += 1
        references=list(w['Public Evidence for SEC FInding'].values)
        ref_header=next((i for i,row in enumerate(references) if row[0]=='Evidence ID'),None)
        if ref_header is None: raise ValueError('Public evidence header not found')
        for row in references[ref_header+1:]:
            vals=[cell_text(x) for x in (list(row[:8])+[None]*8)[:8]]
            if not vals[0]: continue
            conn.execute('''INSERT INTO public_evidence
              (id,related_findings,source,date_vintage,observation,match_timing,limitations,source_url,origin)
              VALUES (?,?,?,?,?,?,?,?,?)
              ON CONFLICT(id) DO UPDATE SET observation=excluded.observation,source_url=excluded.source_url''',
              (*vals,ORIGIN))
            summary['public_evidence'] += 1
        for row in list(w['Files Summary'].values)[1:]:
            vals=[cell_text(x) for x in (list(row[:12])+[None]*12)[:12]]
            if not vals[0]:continue
            conn.execute('''INSERT INTO sources(id,name,description,record_level,facility_matching,
              history,update_schedule,access_method,license_notes,uses,limitations,core_tables,url,origin)
              VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
              ON CONFLICT(id) DO UPDATE SET description=excluded.description,origin=excluded.origin''',
              (slug(vals[0]),*vals,ORIGIN))
            summary['sources'] += 1
        for row in list(w['Field Information'].values)[1:]:
            vals=[cell_text(x) for x in (list(row[:11])+[None]*11)[:11]]
            if not vals[0]:continue
            conn.execute('''INSERT INTO field_catalog
              (id,source,dataset,table_name,purpose,field_count,join_identifiers,key_risk_fields,
              actual_fields,download_url,documentation_url,caveats,origin)
              VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
              ON CONFLICT(id) DO UPDATE SET actual_fields=excluded.actual_fields,origin=excluded.origin''',
              (stable_id('fields',vals[0],vals[1],vals[2]),*vals,ORIGIN))
            summary['catalog_tables'] += 1
        conn.commit()
        summary_dict = dict(summary)
        summary_dict['input_sha256']=sha
        summary_dict['ingestion_run_id']=run_id
        complete_run(conn,run_id,utcnow(),sum(summary.values()),summary_dict)
        return summary_dict
    except Exception as exc:
        conn.rollback()
        complete_run(conn,run_id,utcnow(),0,{'error':str(exc)},status='failed')
        raise
    finally:
        w.close()
