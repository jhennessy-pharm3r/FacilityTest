"""Release-scoped official CSB reporting schema, markers and deterministic output."""
import csv
import gzip
import json
from openpyxl import Workbook
import pytest
from scripts.build_csb_incidents import HEADERS, stage_csb, validate_csb

def fixture_book(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Public"
    sheet.append(["Accidental Release Events"])
    sheet.append(HEADERS)
    sheet.append([1, "2023-07-09", "A reported facility", "Houston", "TX",
                  "Yes", "No", "Yes", "Reported"])
    sheet.append([2, "2023-08-22", "Another company", "Cleveland", "OH",
                  "No", "Yes", "", "Filed"])
    sheet.append([3, "", "Company without date", "Dallas", "TX",
                  "No", "No", "No", "Reported"])
    book.save(path)

def test_csb_release_scoped_ids_and_markers(tmp_path):
    raw = tmp_path/"csb.xlsx"
    fixture_book(raw)
    manifest = tmp_path/"manifest.json"
    out = tmp_path/"official.csv.gz"
    result = stage_csb(raw, out, manifest, release="2026-09-01", minimum_rows=2)
    assert result["raw_data_rows"] == 3
    assert result["accepted_rows"] == 2
    assert result["rejected_rows"] == 1
    assert result["rejection_reasons"] == {"missing_or_invalid_incident_date":1}
    assert result["state_counts"] == {"OH":1,"TX":1}
    assert result["qualitative_marker_counts"]["Fatality"] == {"no_as_reported":1,"yes_as_reported":1}
    assert result["qualitative_marker_counts"]["Report"] == {"other_present_not_interpreted":2}
    assert "Company without date" not in str(result)
    assert validate_csb(out,result)
    with gzip.open(out,"rt") as file:
        rows = list(csv.DictReader(file))
    assert [x["source_id"] for x in rows] == ["csb:2026-09-01:1","csb:2026-09-01:2"]
    assert rows[0]["fatality_marker_as_reported"] == "Yes"
    assert "fatalities" not in rows[0]  # not an actual confirmed count
    assert rows[0]["coordinates"] == ""
    assert all(x["site_link_status"] == "candidate_only_not_matched" for x in rows)
    second = stage_csb(raw,tmp_path/"different-name.csv.gz",tmp_path/"m2.json",
                       release="2026-09-01",minimum_rows=2)
    assert second["normalized_sha256"] == result["normalized_sha256"]
    assert json.loads(manifest.read_text())["source_identity_key"] == "release_date + original Number"

def test_csb_schema_drift_rejected(tmp_path):
    raw = tmp_path/"changed.xlsx"
    fixture_book(raw)
    from openpyxl import load_workbook
    book = load_workbook(raw)
    book["Public"]["B2"] = "Some new date"
    book.save(raw)
    with pytest.raises(ValueError, match="schema drift"):
        stage_csb(raw,tmp_path/"fail.csv.gz",tmp_path/"fail.json",release="2026-09-01",minimum_rows=2)

def test_csb_tamper_detected(tmp_path):
    raw=tmp_path/"input.xlsx"
    fixture_book(raw)
    out=tmp_path/"result.csv.gz"
    report=stage_csb(raw,out,tmp_path/"manifest.json",release="2026-09-01",minimum_rows=2)
    out.write_bytes(out.read_bytes()+b"X")
    with pytest.raises(ValueError, match="checksum"):
        validate_csb(out,report)
