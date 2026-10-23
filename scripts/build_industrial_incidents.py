"""Lane C source-preserving industrial incident intake (no physical-site auto-linking).

Accepts official CSV/XLSX/ZIP-exported tabular datasets supplied as local files.
Raw inputs must be downloaded from the official agencies, archived, and retained.
No unverified NRC report is a confirmed accident or insured loss.
"""
from __future__ import annotations

import argparse
import csv
from datetime import date, datetime, timezone
import gzip
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook
try:
    from scripts.industrial_incident_io import gzip_csv_writer
except ModuleNotFoundError:  # direct python scripts/ entrypoint
    from industrial_incident_io import gzip_csv_writer

PUBLISHERS = {
    "nrc": ("USCG National Response Center", "reported_unverified",
            "https://nrc.uscg.mil/"),
    "csb": ("U.S. Chemical Safety and Hazard Investigation Board", "reported_to_agency",
            "https://www.csb.gov/news/incident-report-rule-form-/"),
    "phmsa_pipeline": ("PHMSA Office of Pipeline Safety", "operator_reported",
            "https://www.phmsa.dot.gov/data-and-statistics/pipeline/source-data"),
    "phmsa_hazmat": ("PHMSA Hazardous Materials Safety", "reporter_submitted",
            "https://www.phmsa.dot.gov/hazmat-program-management-data-and-statistics/data-operations/incident-statistics"),
}
OUT_COLUMNS = ("source","source_id","incident_date","date_precision","name","city","state",
               "latitude","longitude","coordinate_precision","reported_fatalities",
               "reported_injuries","reported_damage_usd","evidence_state",
               "site_link_state","source_url","raw_table","raw_row_number")
ALIASES = {
    "nrc": {
        "id": ("seqnos", "seqno", "sequence_number", "sequence", "nrc_report_number"),
        "date": ("incident_date", "incidentdate", "date_of_incident", "incident_date_time", "incident_dtg"),
        "name": ("responsible_company", "responsible_party", "company_name", "company"),
        "city": ("incident_city", "city", "location_nearest_city"), "state": ("incident_state", "state", "location_state"),
        "lat": ("latitude", "lat"), "lon": ("longitude", "lon"),
        "fatal": ("fatalities", "deaths"), "injury": ("injuries",),
        "damage": ("damage", "property_damage"),
    },
    "csb": {
        "id": ("csb_id", "incident_id", "report_number", "event_id", "number"),
        "date": ("incident_date", "date_of_incident", "date_of_event", "event_date"),
        "name": ("company_name", "facility_name", "company", "facility"),
        "city": ("city", "incident_city"), "state": ("state", "incident_state"),
        "lat": ("latitude", "lat"), "lon": ("longitude", "lon"),
        "fatal": ("fatalities", "number_of_fatalities"), "injury": ("serious_injuries", "injuries"),
        "damage": ("property_damage", "property_damage_usd"),
    },
    "phmsa_pipeline": {
        "id": ("report_number", "report_no", "incident_id", "report_id"),
        "date": ("incident_date", "date_of_incident", "local_datetime"),
        "name": ("operator_name", "operator", "company_name"),
        "city": ("incident_city", "city"), "state": ("incident_state", "state"),
        "lat": ("latitude", "lat"), "lon": ("longitude", "lon"),
        "fatal": ("fatalities", "fatal", "number_of_fatalities"),
        "injury": ("injuries", "injured", "number_of_injuries"),
        "damage": ("total_cost", "property_damage", "total_cost_current"),
    },
    "phmsa_hazmat": {
        "id": ("report_number", "report_no", "incident_id"),
        "date": ("date_of_incident", "incident_date"),
        "name": ("carrier_reporter_name", "carrier_name", "reporter_name"),
        "city": ("incident_city", "city"), "state": ("incident_state", "state"),
        "lat": ("latitude", "lat"), "lon": ("longitude", "lon"),
        "fatal": ("fatalities", "deaths"), "injury": ("injuries",),
        "damage": ("property_damage", "total_cost"),
    },
}

def digest(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def key(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")

def select_headers(headers, source, overrides=None):
    found = {key(h): h for h in headers if h is not None}
    selected = {}
    for label, candidates in ALIASES[source].items():
        override = (overrides or {}).get(label)
        options = (override,) if override else candidates
        selected[label] = next((found[key(x)] for x in options if key(x) in found), None)
    if not selected["id"] or not selected["date"]:
        raise ValueError("Source ID/date columns not resolved; supply --mapping JSON with id/date names; observed: " + repr(list(found)[:50]))
    return selected

def rows_from_file(path, *, sheet=None, member=None, header_row=1, encoding="utf-8-sig"):
    if encoding not in ("utf-8-sig", "cp1252", "latin-1"):
        raise ValueError("Unsupported encoding; require verified publisher text encoding")
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".zip":
        with zipfile.ZipFile(path) as archive:
            names = [i for i in archive.infolist() if not i.is_dir()]
            if any(i.file_size > 600_000_000 or ".." in Path(i.filename).parts or
                   i.filename.startswith("/") for i in names):
                raise ValueError("Unsafe archive member")
            choices = [i for i in names if i.filename.lower().endswith((".csv", ".txt"))]
            if member:
                choices = [i for i in choices if i.filename == member]
            if len(choices) != 1:
                raise ValueError("Choose a single CSV/TXT zip member with --member")
            name = choices[0].filename
            with archive.open(name) as binary:
                with io.TextIOWrapper(binary, encoding=encoding, errors="strict", newline="") as data:
                    dialect = csv.excel_tab if name.lower().endswith(".txt") else csv.excel
                    yield from enumerate(csv.DictReader(data, dialect=dialect), 2)
        return
    if ext in (".csv", ".txt"):
        with path.open("r", encoding=encoding, errors="strict", newline="") as handle:
            dialect = csv.excel_tab if ext == ".txt" else csv.excel
            yield from enumerate(csv.DictReader(handle, dialect=dialect), 2)
        return
    if ext == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if not sheet or sheet not in book.sheetnames:
                raise ValueError("Explicit --sheet required for NRC/CSB XLSX; available sheets: " + repr(book.sheetnames))
            if header_row < 1 or header_row > 50:
                raise ValueError("header_row must be between 1 and 50")
            iterator = book[sheet].iter_rows(min_row=header_row, values_only=True)
            headings = [str(x).strip() if x is not None else "" for x in next(iterator)]
            for number, values in enumerate(iterator, header_row + 1):
                if any(value is not None for value in values):
                    yield number, dict(zip(headings, values))
        finally:
            book.close()
        return
    raise ValueError("Supported source formats: csv, txt, xlsx, zip")

def as_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat(), "day"
    if isinstance(value, date):
        return value.isoformat(), "day"
    if isinstance(value, (int, float)) and 1900 <= int(value) <= 2100:
        return str(int(value)), "year"
    s = str(value or "").strip()
    if not s:
        raise ValueError("missing incident date")
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(s, pattern).date().isoformat(), "day"
        except ValueError:
            continue
    raise ValueError("unrecognized incident date")

def number(value, *, bounds=None):
    s = str(value or "").replace(",", "").replace("$", "").strip()
    if not s:
        return ""
    v = float(s)
    if not math.isfinite(v) or (bounds and not bounds[0] <= v <= bounds[1]) or (not bounds and v < 0):
        raise ValueError("invalid numeric observation")
    return str(v)

def stage(source, raw, normalized, manifest, *, source_url, sheet=None, member=None,
          mapping=None, minimum_rows=1, header_row=1, encoding="utf-8-sig"):
    if source not in PUBLISHERS:
        raise ValueError("Unknown agency source")
    if not source_url.startswith("https://"):
        raise ValueError("Source provenance must be HTTPS")
    raw, normalized, manifest = map(Path, (raw, normalized, manifest))
    normalized.parent.mkdir(parents=True, exist_ok=True)
    manifest.parent.mkdir(parents=True, exist_ok=True)
    temp = Path(str(normalized) + ".part")
    seen = set()
    accepted = 0
    rejected = 0
    duplicate = 0
    missing_coords = 0
    failures = {}
    selected = None
    try:
        with gzip_csv_writer(temp) as target:
            writer = csv.DictWriter(target, fieldnames=OUT_COLUMNS)
            writer.writeheader()
            for row_num, row in rows_from_file(raw, sheet=sheet, member=member, header_row=header_row, encoding=encoding):
                if selected is None:
                    selected = select_headers(row.keys(), source, mapping)
                def get(field):
                    return row.get(selected[field]) if selected[field] else None
                identifier = str(get("id") or "").strip()
                if not identifier:
                    rejected += 1
                    failures["missing_id"] = failures.get("missing_id", 0) + 1
                    continue
                if identifier in seen:
                    duplicate += 1
                    rejected += 1
                    failures["duplicate_source_id"] = failures.get("duplicate_source_id", 0) + 1
                    continue
                seen.add(identifier)
                try:
                    when, precision = as_date(get("date"))
                    try:
                        lat = number(get("lat"), bounds=(-90, 90))
                        lon = number(get("lon"), bounds=(-180, 180))
                    except ValueError:
                        # Invalid publisher geometry is missing rather than a rejected event.
                        lat = lon = ""
                    if not lat or not lon or (float(lat) == 0 and float(lon) == 0):
                        lat = lon = ""
                    if not lat:
                        missing_coords += 1
                    normalized_row = dict(zip(OUT_COLUMNS, (
                        source, identifier, when, precision, str(get("name") or "").strip()[:500],
                        str(get("city") or "").strip()[:200],
                        str(get("state") or "").strip()[:80],
                        lat, lon, "source_reported_unverified" if lat else "missing",
                        number(get("fatal")), number(get("injury")), number(get("damage")),
                        PUBLISHERS[source][1], "candidate_only_not_matched",
                        source_url, sheet or member or raw.name, row_num
                    )))
                    writer.writerow(normalized_row)
                    accepted += 1
                except (ValueError, TypeError) as error:
                    rejected += 1
                    category = str(error)
                    failures[category] = failures.get(category, 0) + 1
        if accepted < minimum_rows:
            raise ValueError(f"Accepted {accepted} rows; required >= {minimum_rows}")
        temp.replace(normalized)
        summary = {
            "schema_version": "1.0", "source": source, "publisher": PUBLISHERS[source][0],
            "source_url": source_url, "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "raw_sha256": digest(raw), "normalized_sha256": digest(normalized),
            "raw_bytes": raw.stat().st_size, "raw_rows": accepted + rejected,
            "accepted_rows": accepted, "rejected_rows": rejected,
            "duplicate_source_ids": duplicate, "missing_coordinates": missing_coords,
            "failure_reasons": failures, "input_table": sheet or member or raw.name,
            "field_mapping": selected, "state": "STAGED",
            "evidence_caveat": "Source incident observations are not verified physical properties, facility losses, insurance claims, or confirmed exposure; all physical-site links require independent review.",
            "geography": "source reported, not independently verified",
            "deployment": "NOT_DEPLOYED",
        }
        manifest.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
        validate(normalized, summary)
        return summary
    finally:
        temp.unlink(missing_ok=True)

def validate(normalized, manifest):
    if digest(normalized) != manifest["normalized_sha256"]:
        raise ValueError("normalized checksum mismatch")
    seen = set()
    count = 0
    with gzip.open(normalized, "rt", encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        if tuple(reader.fieldnames or ()) != OUT_COLUMNS:
            raise ValueError("normalized schema mismatch")
        for row in reader:
            if row["source_id"] in seen:
                raise ValueError("duplicate normalized source ID")
            seen.add(row["source_id"])
            if row["site_link_state"] != "candidate_only_not_matched":
                raise ValueError("Unexpected physical-site linkage")
            count += 1
    if count != manifest["accepted_rows"] or count + manifest["rejected_rows"] != manifest["raw_rows"]:
        raise ValueError("manifest count mismatch")
    return True

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", choices=PUBLISHERS, required=True)
    parser.add_argument("--raw", type=Path, required=True)
    parser.add_argument("--normalized", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--source-url", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--member")
    parser.add_argument("--encoding", choices=("utf-8-sig","cp1252","latin-1"), default="utf-8-sig", help="Explicit reviewed official CSV/TXT encoding; never silently replace invalid bytes")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--mapping", type=Path, help="JSON field-name overrides for publisher drift")
    parser.add_argument("--min-rows", type=int, default=1)
    opts = parser.parse_args()
    mapping = json.loads(opts.mapping.read_text()) if opts.mapping else None
    print(json.dumps(stage(opts.source, opts.raw, opts.normalized, opts.manifest,
        source_url=opts.source_url, sheet=opts.sheet, member=opts.member,
        mapping=mapping, minimum_rows=opts.min_rows, header_row=opts.header_row, encoding=opts.encoding), indent=2))

if __name__ == "__main__":
    main()
