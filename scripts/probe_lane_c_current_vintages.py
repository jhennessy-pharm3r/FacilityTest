"""Check real current-year NRC and CSB quarter-to-quarter identifier persistence.

Official-only urls. Optional archived vintages can be unavailable; not
proof of zero reports. All source rows remain in disposable runner storage.
"""
from __future__ import annotations
from collections import Counter
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import re
import tempfile
from urllib.parse import urlsplit
import requests
from openpyxl import load_workbook
try:
    from scripts.build_nrc_enriched import stage_nrc
    from scripts.build_csb_incidents import HEADERS,stage_csb
except ModuleNotFoundError:
    from build_nrc_enriched import stage_nrc
    from build_csb_incidents import HEADERS,stage_csb

NRC26="https://nrc.uscg.mil/FOIAFiles/CY26.xlsx"
CSB_Q3="https://www.csb.gov/assets/1/6/Copy_of_Accidental_Release_Events_2026-09-01.xlsx"
CSB_Q2="https://www.csb.gov/assets/1/6/Copy_of_Accidental_Release_Events_2026-06-01.xlsx"
CSB_LANDING="https://www.csb.gov/news/incident-report-rule-form-/"

def official_download(url,target,limit=60_000_000):
    approved={"nrc.uscg.mil","www.csb.gov"}
    sha=hashlib.sha256()
    size=0
    with requests.get(url,stream=True,timeout=(15,95),headers={
            "User-Agent":"FacilityRisk-LaneC-2026-OfficialCoverageAudit/1.0",
            "Accept-Encoding":"identity"}) as r:
        r.raise_for_status()
        parsed=urlsplit(r.url)
        if parsed.scheme!="https" or parsed.hostname not in approved:
            raise ValueError("Publisher redirected beyond approved first-party official hosts")
        with target.open("wb") as output:
            for chunk in r.iter_content(1024*1024):
                if not chunk:continue
                size+=len(chunk)
                if size>limit:raise ValueError("Official file exceeds safety cap")
                output.write(chunk);sha.update(chunk)
    if size<1024:raise ValueError("Official source unexpectedly small")
    return {"downloaded_bytes":size,"raw_sha256":sha.hexdigest(),"url":url}

def csb_release_index(path):
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=book["Public"]
        rows=ws.iter_rows(values_only=True)
        next(rows)
        fields=tuple(str(x or "").strip() for x in next(rows)[:9])
        if fields!=HEADERS:
            raise ValueError("Official CSB archived quarter has distinct schema, cannot assume native ID stability")
        out={}
        duplicates=0
        for values in rows:
            if not any(x is not None for x in values):continue
            row=dict(zip(HEADERS,values[:9]))
            source_id=str(row["Number"] or "").strip()
            if not source_id or source_id in out:
                duplicates+=1
                continue
            incident_date=str(row["Incident\nDate"] or "").strip()
            out[source_id]=(
                incident_date,
                str(row["Company"] or "").strip().casefold(),
                str(row["City"] or "").strip().casefold(),
                str(row["State"] or "").strip().casefold()
            )
        if duplicates:raise ValueError("CSB historical release contains duplicate or missing native Number values")
        return out
    finally:book.close()

def compare_quarters(old,current):
    """Publisher Number only; do not declare identity if any event attributes differ."""
    old_ids=set(old)
    new_ids=set(current)
    overlap=old_ids & new_ids
    identical=sum(old[i]==current[i] for i in overlap)
    changed=len(overlap)-identical
    return {
        "old_release_rows":len(old),"new_release_rows":len(current),
        "reused_number_count":len(overlap),
        "unchanged_four_field_event_identity_for_same_number":identical,
        "changed_event_attributes_under_same_number":changed,
        "numbers_absent_in_new_release":len(old_ids-new_ids),
        "new_number_values":len(new_ids-old_ids),
        "safe_cross_vintage_id_assumption":False,
        "note":"Only same Number and exactly same source event date/company/city/state provides provisional continuity. It does not establish permanent agency native ID or physical site."
    }

def main(output):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    audit={"checked_at":datetime.now(timezone.utc).isoformat(),
           "geographic_or_physical_site_matches":0,"deployed":0,
           "current_year_nrc":{"state":"NOT_TESTED","vintage":"2026","official_url":NRC26},
           "csb_current_release":{"state":"NOT_TESTED","url":CSB_Q3},
           "csb_prior_quarter":{"state":"NOT_TESTED","url":CSB_Q2}}
    try:
        r=requests.get(CSB_LANDING,timeout=(12,25),headers={
            "User-Agent":"FacilityRisk-LaneC-CatalogCheck/1.0"})
        r.raise_for_status()
        if urlsplit(r.url).hostname!="www.csb.gov":
            raise ValueError("CSB publisher redirected to unexpected website")
        latest=sorted(set(re.findall(r"Copy_of_Accidental_Release_Events_(20\\d{2}-\\d{2}-\\d{2})\\.xlsx",r.text)))
        audit["csb_release_catalog"]={"state":"LIVE_OFFICIAL_PAGE_INVENTORIED",
             "latest_filename_date":max(latest) if latest else None,
             "discovered_vintages_in_page":latest,
             "page_url":CSB_LANDING}
    except Exception as e:
        audit["csb_release_catalog"]={"state":"BLOCKED","error":type(e).__name__+": "+str(e)[:180]}
    with tempfile.TemporaryDirectory() as t:
        wd=Path(t)
        try:
            path=wd/"CY26.xlsx"
            receipt=official_download(NRC26,path)
            result=stage_nrc(path,wd/"nrc26.gz",wd/"nrc26.json",year=2026,minimum_rows=1)
            audit["current_year_nrc"].update(receipt)
            audit["current_year_nrc"].update({
                "state":"PARTIAL_CURRENT_CALENDAR_YEAR_MASTER_ROWS_STAGED_NOT_INDEXED",
                "accepted":result["accepted_rows"],
                "missing_child_links":result["missing_cross_table_links"],
                "coordinates":result["coordinate_records"],
                "report_is_not_verified_incident":True})
        except Exception as e:
            audit["current_year_nrc"].update({"state":"OFFICIAL_PUBLISHER_UNAVAILABLE_OR_SCHEMA_BLOCKED",
                "error":type(e).__name__+": "+str(e)[:300],
                "accepted":None})
        q={}
        for key,url in (("csb_current_release",CSB_Q3),("csb_prior_quarter",CSB_Q2)):
            try:
                path=wd/(key+".xlsx")
                q[key]=csb_release_index(path) if path.exists() else None
                receipt=official_download(url,path,30_000_000)
                q[key]=csb_release_index(path)
                audit[key].update(receipt)
                audit[key]["state"]="DOWNLOADED_AND_SCHEMA_VALIDATED_RELEASE_ONLY"
                audit[key]["release_rows_with_unique_number"]=len(q[key])
            except Exception as e:
                audit[key].update({"state":"UNAVAILABLE_OR_SCHEMA_BLOCKED",
                    "error":type(e).__name__+": "+str(e)[:300],
                    "release_rows_with_unique_number":None})
        if q.get("csb_prior_quarter") is not None and q.get("csb_current_release") is not None:
            audit["csb_cross_release_identity_review"]=compare_quarters(
                q["csb_prior_quarter"],q["csb_current_release"])
        else:
            audit["csb_cross_release_identity_review"]={"state":"NOT_TESTED_NO_OFFICIAL_PRIOR_RELEASE"}
    (output/"lane-c-current-vintage-audit.json").write_text(json.dumps(audit,indent=2)+"\n")
    print(json.dumps(audit,indent=2))
    return audit

if __name__=="__main__":
    import argparse
    p=argparse.ArgumentParser()
    p.add_argument("--output",type=Path,required=True)
    opts=p.parse_args()
    main(opts.output)
