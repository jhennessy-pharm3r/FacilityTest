"""Inspect official USCG National Response Center annual Excel data safely.

Probing workbook reachability/layout is NOT incident ingestion or facility
matching. Reports are initially unverified; never promote them to claims.
"""
from __future__ import annotations

import argparse
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import re
import sys
import time
from urllib.parse import urlsplit
import zipfile

import requests
from openpyxl import load_workbook

MAX_COMPRESSED = 50 * 1024 * 1024
MAX_UNCOMPRESSED = 800 * 1024 * 1024
HEADER_TOKENS = re.compile(r"(?i)(?:incident|sequence|report|date|time|state|county|"
    r"location|company|material|casualty|spill|discharge|quantity|cause|fatal|"
    r"injur|evacuation|record|nrc|seqno|type|release)")


def official_url(year):
    year=int(year)
    if not 1990<=year<=2100:
        raise ValueError("NRC annual workbook year outside known range")
    return f"https://nrc.uscg.mil/FOIAFiles/CY{year%100:02d}.xlsx"


def sha256(path):
    result=hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b""):
            result.update(chunk)
    return result.hexdigest()


def download(session,url,dest):
    if urlsplit(url).scheme!="https" or urlsplit(url).hostname!="nrc.uscg.mil" or \
            not re.fullmatch(r"/FOIAFiles/CY\d{2}\.xlsx",urlsplit(url).path):
        raise ValueError("Only official NRC annual workbook URLs are permitted")
    dest=Path(dest)
    dest.parent.mkdir(parents=True,exist_ok=True)
    temp=Path(str(dest)+".part")
    for attempt in range(4):
        try:
            with session.get(url,stream=True,timeout=(20,90),
                headers={"User-Agent":"Pharm3r-FacilityRisk-NRC-Probe/1.0",
                         "Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         "Accept-Encoding":"identity"}) as response:
                response.raise_for_status()
                redirect=urlsplit(response.url)
                if redirect.scheme!="https" or redirect.hostname!="nrc.uscg.mil":
                    raise ValueError("NRC data redirected to unapproved publisher")
                size=0
                with temp.open("wb") as handle:
                    for chunk in response.iter_content(chunk_size=1024*1024):
                        if chunk:
                            size+=len(chunk)
                            if size>MAX_COMPRESSED:
                                raise ValueError("NRC workbook unexpectedly large")
                            handle.write(chunk)
                if size<4096 or temp.open("rb").read(2)!=b"PK":
                    raise ValueError("NRC response is not a substantive XLSX archive")
            temp.replace(dest)
            return dest
        except requests.RequestException as exc:
            status=getattr(getattr(exc,"response",None),"status_code",None)
            if attempt==3 or (status is not None and status!=429 and status<500):
                raise
            pause=2**attempt
            print(f"NRC transient download error; retry {attempt+1}/4 in {pause}s",
                  file=sys.stderr,flush=True)
            time.sleep(pause)
        finally:
            temp.unlink(missing_ok=True)
    raise RuntimeError("Unreachable NRC retry boundary")


def probe(path,year,min_sheets=2):
    file=Path(path)
    with zipfile.ZipFile(file) as archive:
        info=archive.infolist()
        total=sum(z.file_size for z in info)
        if total>MAX_UNCOMPRESSED or any(
            z.filename.startswith("/") or ".." in Path(z.filename).parts
            for z in info
        ):
            raise ValueError("Unsafe NRC workbook archive members")
        if "xl/workbook.xml" not in {z.filename for z in info}:
            raise ValueError("NRC publisher file not an Excel workbook")
    book=load_workbook(file,read_only=True,data_only=True)
    try:
        if len(book.worksheets)<min_sheets:
            raise ValueError("NRC annual workbook missing expected multiple tables")
        tables=[]
        for sheet in book.worksheets:
            # Public field names are safe; never print/sample underlying
            # incident rows, company names, private telephone numbers, etc.
            first=next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())
            titles=[str(x).strip() if isinstance(x,str) else "" for x in first]
            supported=[x for x in titles if HEADER_TOKENS.search(x)]
            observed_header=(len(supported)>=2 and
                             len(supported)>=max(2,len([x for x in titles if x])//5))
            headers=titles[:80] if observed_header else []
            tables.append({
                "sheet":sheet.title[:100],
                "excel_reported_rows":sheet.max_row,
                "excel_reported_columns":sheet.max_column,
                "header_recognized":observed_header,
                "header_names":headers,
            })
    finally:
        book.close()
    count=sum(1 for s in tables if s["header_recognized"])
    if count<1:
        raise ValueError("No recognizable NRC incident/master relationship headers")
    return {"schema_version":"1.0","source":"USCG NRC annual FOIA workbook",
            "state":"schema_probe_only_not_imported",
            "year":year,"source_url":official_url(year),
            "retrieved_at":datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "sha256":sha256(file),"archive_bytes":file.stat().st_size,
            "archive_expanded_bytes":total,"worksheets":tables,
            "recognizable_headers":count,
            "data_caveat":"NRC notifications are initially unverified; sheet row counts are not distinct incident counts and no facility matching/claims are asserted"}


def main():
    p=argparse.ArgumentParser(description="Examine actual Coast Guard NRC annual workbook layout")
    p.add_argument("--year",type=int,required=True)
    p.add_argument("--output",type=Path,required=True)
    p.add_argument("--report",type=Path,required=True)
    args=p.parse_args()
    url=official_url(args.year)
    download(requests.Session(),url,args.output)
    result=probe(args.output,args.year)
    args.report.parent.mkdir(parents=True,exist_ok=True)
    args.report.write_text(json.dumps(result,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"source":url,"schema_probe":result["state"],
        "year":args.year,"workbook_bytes":result["archive_bytes"],
        "sheets":[{"name":r["sheet"],"rows_reported":r["excel_reported_rows"],
                   "columns":r["excel_reported_columns"],"headers":r["header_names"]}
                  for r in result["worksheets"]]},indent=2),flush=True)


if __name__=="__main__":
    main()
