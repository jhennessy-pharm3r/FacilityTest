"""Probe PHMSA-derived nationwide annual trends released by DOT BTS (Table 2-50).

These are AGGREGATED annual statistics, not pipeline event records, individual
facilities, insured losses, or a substitute for official PHMSA operator ZIPs.
"""
from __future__ import annotations
import argparse
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import tempfile
from urllib.parse import urlsplit
import requests
from openpyxl import load_workbook

URL="https://www.bts.gov/sites/bts.dot.gov/files/2026-02/table_02_50_022526.xlsx"
PUBLISHER="US DOT Bureau of Transportation Statistics / PHMSA-derived National Transportation Statistics Table 2-50"
def fetch(dest):
    with requests.get(URL,stream=True,timeout=(20,90),headers={
        "User-Agent":"FacilityRisk-Industrial-LaneC/1.0","Accept-Encoding":"identity"}) as r:
        r.raise_for_status()
        parsed=urlsplit(r.url)
        if parsed.scheme!="https" or parsed.hostname!="www.bts.gov":
            raise ValueError("BTS official source redirect outside approved publisher")
        sha=hashlib.sha256()
        size=0
        with dest.open("wb") as output:
            for chunk in r.iter_content(1024*1024):
                if not chunk:continue
                size+=len(chunk)
                if size>4_000_000:raise ValueError("Unexpectedly large official BTS XLSX")
                output.write(chunk)
                sha.update(chunk)
    return size,sha.hexdigest()
def probe(path):
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        tables=[]
        for sheet in book.worksheets:
            it=sheet.iter_rows(values_only=True)
            head=[]
            rows=0
            year_columns=[]
            for i,line in enumerate(it,1):
                if not any(x is not None for x in line):continue
                rows+=1
                if len(head)<16:
                    # Source is national aggregate statistical table only.
                    head.append({"excel_row":i, "cells":[
                        str(x or "")[:90] for x in line[:16]]})
                first=line[0] if line else None
                if isinstance(first,(int,float)) and 1970<=first<=2026:
                    year_columns.append(int(first))
            tables.append({"sheet":sheet.title,"nonempty_rows":rows,
                           "first_rows":head,
                           "observed_numeric_first_column_years":year_columns[:7],
                           "year_like_first_column_count":len(year_columns)})
        return tables
    finally:
        book.close()
def main():
    p=argparse.ArgumentParser()
    p.add_argument("--output",type=Path,required=True)
    args=p.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    audit={"source":PUBLISHER,"source_url":URL,"retrieved_at":datetime.now(timezone.utc).isoformat(),
           "record_grain":"national/yearly aggregate; never individual PHMSA accidents",
           "state":"NOT_TESTED","event_records_downloaded":0,
           "production_deployed":0}
    with tempfile.TemporaryDirectory() as dir:
        file=Path(dir)/"table_02_50_022526.xlsx"
        try:
            audit["raw_bytes"],audit["raw_sha256"]=fetch(file)
            audit["tables"]=probe(file)
            audit["state"]="OFFICIAL_AGGREGATE_DOWNLOADED_SCHEMA_INVENTORIED_NOT_NORMALIZED"
        except Exception as exc:
            audit["state"]="BLOCKED"
            audit["error"]=type(exc).__name__+": "+str(exc)[:250]
        (args.output/"bts-phmsa-official-probe.json").write_text(json.dumps(audit,indent=2)+"\n")
    print(json.dumps(audit,indent=2))
    if audit["state"]=="BLOCKED":raise SystemExit(1)
if __name__=="__main__":
    main()
