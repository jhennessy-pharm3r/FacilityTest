"""Source-faithful CSB 40 CFR Part 1604 quarterly release ingestion.

CSB's 'Number' is a row ID within this published workbook, NOT proven stable
across later releases. Fatality, injury and damage columns are reported
qualitative markers, never inferred casualty counts or monetary losses.
"""
from __future__ import annotations
from collections import Counter
import csv
from datetime import datetime, timezone
import gzip
import json
from pathlib import Path
from openpyxl import load_workbook

try:
    from scripts.build_industrial_incidents import as_date, digest
    from scripts.industrial_incident_io import gzip_csv_writer
except ModuleNotFoundError:  # direct scripts/probe entrypoint
    from build_industrial_incidents import as_date, digest
    from industrial_incident_io import gzip_csv_writer

HEADERS = ("Number", "Incident\nDate", "Company", "City", "State",
           "Fatality", "Serious\nInjury", "Substantial Property Damage", "Report")
COLUMNS = ("source", "source_id", "source_release", "source_row_number",
           "original_number", "incident_date", "date_precision", "company_as_reported",
           "incident_city", "incident_state", "fatality_marker_as_reported",
           "serious_injury_marker_as_reported", "substantial_damage_marker_as_reported",
           "report_marker_as_reported", "coordinates", "evidence_status",
           "site_link_status", "source_document")

SOURCE = "https://www.csb.gov/assets/1/6/Copy_of_Accidental_Release_Events_2026-09-01.xlsx"

def _marker(value):
    # Publisher markers are not counts. Never map Yes/N/A/blank to zero or one.
    return str(value).strip()[:120] if value is not None else ""

def stage_csb(raw, output, manifest, *, release, source_url=SOURCE, minimum_rows=50):
    if not release or len(release) != 10:
        raise ValueError("Explicit CSB source publication vintage YYYY-MM-DD required")
    raw, output, manifest = map(Path, (raw, output, manifest))
    output.parent.mkdir(parents=True, exist_ok=True)
    manifest.parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(str(output) + ".part")
    seen = set()
    states, markers, problems = Counter(), {key: Counter() for key in HEADERS[5:]}, Counter()
    accepted, rejected, blank_city = 0, 0, 0
    dates = []
    book = load_workbook(raw, read_only=True, data_only=True)
    try:
        if "Public" not in book.sheetnames:
            raise ValueError("CSB official workbook Public sheet missing")
        values = book["Public"].iter_rows(values_only=True)
        next(values)  # publisher title row, not actual column header
        header = tuple(str(v or "").strip() for v in next(values)[:9])
        if header != HEADERS:
            raise ValueError("CSB public schema drift; first nine field names: " + repr(header))
        with gzip_csv_writer(tmp) as handle:
            writer = csv.DictWriter(handle, COLUMNS)
            writer.writeheader()
            for rowno, raw_row in enumerate(values, 3):
                if not any(x is not None for x in raw_row):
                    continue
                row = dict(zip(HEADERS, raw_row[:9]))
                original_id = _marker(row["Number"])
                if not original_id:
                    problems["missing_original_number"] += 1
                    rejected += 1
                    continue
                if original_id in seen:
                    problems["duplicate_number_within_release"] += 1
                    rejected += 1
                    continue
                seen.add(original_id)
                try:
                    when, date_precision = as_date(row["Incident\nDate"])
                except (ValueError, TypeError):
                    problems["missing_or_invalid_incident_date"] += 1
                    rejected += 1
                    continue
                location = _marker(row["City"])
                state = _marker(row["State"])
                blank_city += not bool(location)
                states[state or "<missing>"] += 1
                for key in markers:
                    markers[key][_marker(row[key]) or "<missing>"] += 1
                writer.writerow(dict(zip(COLUMNS, (
                    "csb_release_events", f"csb:{release}:{original_id}",
                    release, rowno, original_id, when, date_precision,
                    _marker(row["Company"]), location, state,
                    _marker(row["Fatality"]), _marker(row["Serious\nInjury"]),
                    _marker(row["Substantial Property Damage"]), _marker(row["Report"]),
                    "", "csb_reported_accidental_release_unverified_outcome",
                    "candidate_only_not_matched", source_url
                ))))
                dates.append(when)
                accepted += 1
        if accepted < minimum_rows:
            raise ValueError("CSB official cohort unexpectedly small")
        tmp.replace(output)
        report = {
            "source": "CSB accidental release reporting 40 CFR Part 1604",
            "source_url": source_url, "release": release, "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "raw_sha256": digest(raw), "normalized_sha256": digest(output),
            "raw_bytes": raw.stat().st_size, "raw_data_rows": accepted + rejected,
            "accepted_rows": accepted, "rejected_rows": rejected,
            "rejection_reasons": dict(sorted(problems.items())),
            "source_number_scope": "unique within this release only; cross-vintage permanence unverified",
            "state_counts": dict(sorted(states.items())),
            "missing_city": blank_city,
            "first_incident_date": min(dates), "last_incident_date": max(dates),
            "qualitative_marker_counts": {key: dict(sorted(value.items())) for key, value in markers.items()},
            "coordinate_rows": 0,
            "evidence_status": "reported_accidental_release_not_corrob_insured_loss",
            "source_identity_key": "release_date + original Number",
            "indexed_records": 0, "deployed_records": 0, "state": "NORMALIZED_NOT_INDEXED"
        }
        manifest.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        validate_csb(output, report)
        return report
    finally:
        book.close()
        tmp.unlink(missing_ok=True)

def validate_csb(path, manifest):
    if digest(path) != manifest["normalized_sha256"]:
        raise ValueError("CSB checksum mismatch")
    n = 0
    seen = set()
    with gzip.open(path, "rt", encoding="utf-8", newline="") as stream:
        reader = csv.DictReader(stream)
        if tuple(reader.fieldnames or ()) != COLUMNS:
            raise ValueError("CSB normalized schema drift")
        for row in reader:
            if row["source_id"] in seen or not row["source_id"].startswith("csb:" + manifest["release"] + ":"):
                raise ValueError("CSB duplicate or mismatched scoped source ID")
            seen.add(row["source_id"])
            if row["site_link_status"] != "candidate_only_not_matched" or row["coordinates"]:
                raise ValueError("CSB location wrongly promoted to a verified property point")
            n += 1
    if n != manifest["accepted_rows"] or n + manifest["rejected_rows"] != manifest["raw_data_rows"]:
        raise ValueError("CSB row reconciliation failed")
    return True
